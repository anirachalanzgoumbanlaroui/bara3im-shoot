"""
Base Excel Exporter class providing openpyxl workbook management,
standardized sheet headers, table formatting, cell styling, auto-filtering,
number formatting, summary totals, and HTTP response generation.
"""

from io import BytesIO
from typing import List, Dict, Any, Optional
from django.http import HttpResponse
import openpyxl

from .styles import (
    FONT_TITLE, FONT_SUBTITLE, FONT_SECTION_HEADER,
    FONT_HEADER, FONT_HEADER_ARDIS, FONT_HEADER_SABLLET, FONT_HEADER_GOLD,
    FILL_HEADER_DEFAULT, FILL_HEADER_ARDIS, FILL_HEADER_SABLLET, FILL_HEADER_GOLD,
    FONT_BODY, FONT_BODY_BOLD, FONT_TOTAL, FILL_ZEBRA, FILL_TOTAL,
    ALIGN_LEFT, ALIGN_CENTER, ALIGN_RIGHT,
    ALIGN_HEADER_LEFT, ALIGN_HEADER_CENTER, ALIGN_HEADER_RIGHT,
    BORDER_CELL, BORDER_TOTAL,
    FMT_CURRENCY, FMT_INTEGER, FMT_DECIMAL, FMT_PERCENTAGE, FMT_DATE, FMT_TIME,
    FILL_GREEN, FONT_GREEN, FILL_RED, FONT_RED, FILL_AMBER, FONT_AMBER
)
from .utils import sanitize_sheet_name, generate_export_filename, auto_fit_columns


class BaseExcelExporter:
    """
    Abstract Base Class for all Bara3im Shoot Excel Workbooks.
    """

    def __init__(self, filename: str = "Export.xlsx", theme: str = "default"):
        self.workbook = openpyxl.Workbook()
        # Remove default empty sheet when adding custom sheets
        self._first_sheet_created = False
        self.existing_sheet_names = set()
        self.filename = filename
        self.theme = theme

    def get_header_style(self, theme: str = None):
        """Returns header font and fill based on location/theme."""
        selected_theme = theme or self.theme
        if selected_theme == "ardis":
            return FONT_HEADER_ARDIS, FILL_HEADER_ARDIS
        elif selected_theme == "sabllet":
            return FONT_HEADER_SABLLET, FILL_HEADER_SABLLET
        elif selected_theme == "gold":
            return FONT_HEADER_GOLD, FILL_HEADER_GOLD
        return FONT_HEADER, FILL_HEADER_DEFAULT

    def create_sheet(self, title: str):
        """Creates or reuses a worksheet with sanitized title."""
        clean_title = sanitize_sheet_name(title, self.existing_sheet_names)
        if not self._first_sheet_created:
            ws = self.workbook.active
            ws.title = clean_title
            self._first_sheet_created = True
        else:
            ws = self.workbook.create_sheet(title=clean_title)

        ws.views.sheetView[0].showGridLines = True
        return ws

    def add_header_block(self, ws, title: str, subtitle: str = None, metadata: Dict[str, Any] = None) -> int:
        """
        Writes a clean executive title block at top of sheet.
        Returns the next available row index.
        """
        current_row = 1
        ws.cell(row=current_row, column=1, value=title).font = FONT_TITLE
        current_row += 1

        if subtitle:
            ws.cell(row=current_row, column=1, value=subtitle).font = FONT_SUBTITLE
            current_row += 1

        if metadata:
            meta_str = " | ".join([f"{k}: {v}" for k, v in metadata.items() if v is not None])
            if meta_str:
                ws.cell(row=current_row, column=1, value=meta_str).font = FONT_SUBTITLE
                current_row += 1

        current_row += 1  # Blank spacing line
        return current_row

    def write_table(
        self,
        ws,
        start_row: int,
        columns: List[Dict[str, Any]],
        data: List[Dict[str, Any]],
        theme: str = None,
        include_totals: bool = False,
        total_label_col: str = None,
    ) -> int:
        """
        Writes a structured table with headers, formatted rows, alternate zebra fill,
        proper alignments, and summary total rows with Excel formulas.

        Column definition structure:
        {
            'key': 'field_name',
            'label': 'Header Text',
            'align': 'left'|'center'|'right',
            'format': 'currency'|'integer'|'decimal'|'percentage'|'date'|'time'|'status'|'text',
            'formula_total': 'SUM'|'AVERAGE'|None
        }
        """
        header_font, header_fill = self.get_header_style(theme)
        current_row = start_row

        # Write Header Row
        ws.row_dimensions[current_row].height = 28
        for col_idx, col_def in enumerate(columns, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=col_def['label'])
            cell.font = header_font
            cell.fill = header_fill
            cell.border = BORDER_CELL
            align_type = col_def.get('align', 'left')
            if align_type == 'center':
                cell.alignment = ALIGN_HEADER_CENTER
            elif align_type == 'right':
                cell.alignment = ALIGN_HEADER_RIGHT
            else:
                cell.alignment = ALIGN_HEADER_LEFT

        header_row_idx = current_row
        current_row += 1
        data_start_row = current_row

        # Write Data Rows
        for r_idx, row_data in enumerate(data):
            ws.row_dimensions[current_row].height = 20
            use_zebra = (r_idx % 2 == 1)

            for col_idx, col_def in enumerate(columns, 1):
                key = col_def['key']
                val = row_data.get(key, '')
                cell = ws.cell(row=current_row, column=col_idx)

                align_type = col_def.get('align', 'left')
                fmt_type = col_def.get('format', 'text')

                if align_type == 'center':
                    cell.alignment = ALIGN_CENTER
                elif align_type == 'right':
                    cell.alignment = ALIGN_RIGHT
                else:
                    cell.alignment = ALIGN_LEFT

                cell.border = BORDER_CELL
                if use_zebra:
                    cell.fill = FILL_ZEBRA

                cell.font = FONT_BODY

                # Value formatting
                if fmt_type == 'currency':
                    try:
                        cell.value = float(val) if val is not None else 0.0
                    except (ValueError, TypeError):
                        cell.value = 0.0
                    cell.number_format = FMT_CURRENCY
                elif fmt_type == 'integer':
                    try:
                        cell.value = int(val) if val is not None else 0
                    except (ValueError, TypeError):
                        cell.value = 0
                    cell.number_format = FMT_INTEGER
                elif fmt_type == 'decimal':
                    try:
                        cell.value = float(val) if val is not None else 0.0
                    except (ValueError, TypeError):
                        cell.value = 0.0
                    cell.number_format = FMT_DECIMAL
                elif fmt_type == 'percentage':
                    try:
                        cell.value = float(val) / 100.0 if float(val) > 1 else float(val) if val is not None else 0.0
                    except (ValueError, TypeError):
                        cell.value = 0.0
                    cell.number_format = FMT_PERCENTAGE
                elif fmt_type == 'date':
                    cell.value = str(val) if val else ''
                    cell.number_format = FMT_DATE
                elif fmt_type == 'time':
                    cell.value = str(val) if val else ''
                    cell.number_format = FMT_TIME
                elif fmt_type == 'status':
                    cell.value = str(val).upper() if val else ''
                    cell.alignment = ALIGN_CENTER
                    status_str = str(val).lower()
                    if status_str in ['present', 'active', 'completed', 'a++', 'a+', 'a', 'high']:
                        cell.fill = FILL_GREEN
                        cell.font = FONT_GREEN
                    elif status_str in ['absent', 'inactive', 'locked', 'd', 'e', 'f', 'low']:
                        cell.fill = FILL_RED
                        cell.font = FONT_RED
                    elif status_str in ['late', 'draft', 'in_progress', 'b', 'c', 'normal']:
                        cell.fill = FILL_AMBER
                        cell.font = FONT_AMBER
                else:
                    cell.value = str(val) if val is not None else ''

            current_row += 1

        data_end_row = current_row - 1

        # Enable AutoFilter on header row
        if data:
            last_col_letter = openpyxl.utils.get_column_letter(len(columns))
            ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{data_end_row}"

        # Freeze Header Row
        ws.freeze_panes = ws[f"A{header_row_idx + 1}"]

        # Summary / Totals Row
        if include_totals and data:
            ws.row_dimensions[current_row].height = 24

            for col_idx, col_def in enumerate(columns, 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = FONT_TOTAL
                cell.fill = FILL_TOTAL
                cell.border = BORDER_TOTAL

                col_key = col_def['key']
                col_letter = openpyxl.utils.get_column_letter(col_idx)
                formula_op = col_def.get('formula_total')
                fmt_type = col_def.get('format', 'text')

                if total_label_col and col_key == total_label_col:
                    cell.value = "TOTAL"
                    cell.alignment = ALIGN_LEFT
                elif formula_op == 'SUM':
                    cell.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{data_end_row})"
                    cell.alignment = ALIGN_RIGHT
                    if fmt_type == 'currency':
                        cell.number_format = FMT_CURRENCY
                    else:
                        cell.number_format = FMT_INTEGER
                elif formula_op == 'AVERAGE':
                    cell.value = f"=AVERAGE({col_letter}{data_start_row}:{col_letter}{data_end_row})"
                    cell.alignment = ALIGN_RIGHT
                    if fmt_type == 'currency':
                        cell.number_format = FMT_CURRENCY
                    elif fmt_type == 'percentage':
                        cell.number_format = FMT_PERCENTAGE
                    else:
                        cell.number_format = FMT_DECIMAL
                else:
                    cell.value = ""

            current_row += 1

        auto_fit_columns(ws)
        return current_row + 1

    def generate_workbook_bytes(self) -> bytes:
        """Saves openpyxl workbook into memory buffer and returns raw bytes."""
        buffer = BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def export_to_response(self, filename: str = None) -> HttpResponse:
        """
        Builds a Django HTTP response stream for file download.
        """
        export_filename = filename or self.filename
        content = self.generate_workbook_bytes()

        response = HttpResponse(
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{export_filename}"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'
        return response
